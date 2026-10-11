"""
Tests del léxico ancashino (solo texto, sin IO externo salvo tmp dirs).

Cubre:
  - canonical_form: variante con otra puntuación/caso -> grafía del corpus
  - gloss_hint: glosa exacta y por subcadena (hook RAG)
  - load_lexicon: ruta inexistente -> None
  - ingest_ancash_lexicon.main(): Excel mínimo -> JSON con entries/variant_index
"""

import importlib.util
import json
from pathlib import Path

import openpyxl
import pytest

from modules.voice.services import text_normalizer_ancash as norm
from modules.voice.services.text_normalizer_ancash import (
    canonical_form,
    gloss_hint,
    load_lexicon,
)


def _load_ingest_module():
    path = Path(__file__).resolve().parents[2] / "scripts" / "ingest_ancash_lexicon.py"
    spec = importlib.util.spec_from_file_location("ingest_ancash_lexicon", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_canonical_form_usa_grafia_del_corpus():
    assert canonical_form("atuqmi ashkallata mikun") == "Atuqmi ashkallata mikun."


def test_canonical_form_desconocido_no_op():
    assert canonical_form("frase que no existe en el corpus xyz") == "frase que no existe en el corpus xyz"


def test_load_lexicon_ruta_inexistente_none(tmp_path):
    assert load_lexicon(tmp_path / "noexiste.json") is None


def test_gloss_hint_exacta():
    hints = gloss_hint("Atuqmi ashkallata mikun.")
    assert hints and hints[0]["es"] == "El zorro come al cordero."
    assert hints[0]["variedad"] == "sihuas"


def test_gloss_hint_subcadena_y_vacia():
    hints = gloss_hint("ashkallata")
    assert 1 <= len(hints) <= 5
    assert all("es" in h for h in hints)
    assert gloss_hint("zzz sin coincidencia qqq") == []
    assert gloss_hint("ab") == []


def test_ingest_excel_minimo(tmp_path):
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quechua Sihuas"
    ws.append(["N", "Variedad", "Oración_Quechua", "Oración_Español"])
    ws.append([1, "Sihuas", "Misim nanan.", "Duele la pierna."])  # dup normalizado abajo
    ws.append([2, "Sihuas", "  misim nanan ", "Duele la pierna."])
    ws.append([3, "Sihuas", "Solo quechua", ""])
    ws.append(["Etiq", None, None, None])
    wb.save(corpus / "QuechuaSihuas.xlsx")

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Quechua Chiquian"
    ws2.append(["N", "Variedad", "Oración_Quechua", "Oración_Español"])
    ws2.append([1, "Chiquian", "Allqu mikun.", "El perro come."])
    ws2.append(["Etiq", None, None, None])
    ent = wb2.create_sheet("Etiquetado")
    ent.append([None, 1, "Huaraz llaqtam.", "", "Huaraz - ciudad"])
    wb2.save(corpus / "QuechuaChiquian.xlsx")

    mod = _load_ingest_module()
    out = tmp_path / "lex.json"
    lex = mod.main(corpus, out)

    assert len(lex["entries"]) == 2  # 1 dup + 1 vacío fuera
    assert lex["meta"]["counts"]["pairs"] == 2
    assert out.exists()
    back = json.loads(out.read_text(encoding="utf-8"))
    assert back["entries"][0]["variedad"] in ("sihuas", "chiquian")
